"""End-to-end test with synthetic plastomes (python tests/test_demo.py)."""
import os, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))
from make_demo_data import build
from micromine.cli import run
from micromine.core import exclusivity, relative_status, variation_type

with tempfile.TemporaryDirectory() as d:
    build(os.path.join(d, "in"))
    genomes, clusters, _ = run(os.path.join(d, "in"), os.path.join(d, "out"), log=lambda *_: None)
    tags = lambda c: ",".join(genomes[g].tag for g in c.genomes)
    by = {c.canon + str(c.ref.start): c for c in clusters}
    acc = [c for c in clusters if c.canon == "ACC"][0]
    assert tags(acc) == "01,02,04", tags(acc)                       # SNP broke the SSR in 03
    uniq = [c for c in clusters if tags(c) == "03"]
    assert len(uniq) == 1 and uniq[0].canon == "AAGTCT"             # exclusive to 03
    at = [c for c in clusters if c.canon == "AT"][0]
    assert {m.reps for m in at.members} == {9, 7}                   # length polymorphism
    assert all(len(c.genomes) == 4 for c in clusters if c is not acc and c not in uniq)
    # higher threshold: AT=7 in 02 falls below 8 and must be recovered as below-threshold
    g2, c2, _ = run(os.path.join(d, "in"), os.path.join(d, "out2"), {2: 8}, log=lambda *_: None)
    at2 = [c for c in c2 if c.canon == "AT"][0]
    assert at2.sub[1].reps == 7 and 1 not in at2.genomes
    assert any(exclusivity(c, 4) == "real" for c in c2)              # exclusive to 03 continua real
    assert variation_type(at, 4) == "Length polymorphism"
    assert variation_type(uniq[0], 4) == "Absent in some species"
    assert relative_status(at, 4)[1][0] == "Shorter"                  # AT=7 in 02 vs. 9 in the others
    # FASTA layouts: per species, per locus (oriented alike, SSR highlighted) and consolidated
    out = os.path.join(d, "out")
    per_sp = os.listdir(os.path.join(out, "fasta_per_species")); assert len(per_sp) == 4
    per_loc = os.listdir(os.path.join(out, "fasta_per_locus")); assert len(per_loc) == len(clusters)
    import re
    for c in clusters:
        recs = open(os.path.join(out, "fasta_per_locus", f"{c.id}_{c.cls}_{c.canon}.fasta")).read().split(">")[1:]
        seqs = {r.split("|")[1][:2]: "".join(r.split("\n")[1:]) for r in recs}
        if "01" in seqs and "04" in seqs:                       # 04 is rotated + reverse strand
            a4, b4 = seqs["01"].upper(), seqs["04"].upper()      # same orientation; the SSR window may be
            assert any(a4[i:] == b4[:len(b4) - i] or b4[i:] == a4[:len(a4) - i]    # shifted by < 1 motif
                       for i in range(0, 7)), c.id
        for r in recs:
            body = "".join(r.split("\n")[1:]); a, b = map(int, re.search(r"ssr_in_seq=(\d+)-(\d+)", r).groups())
            ssr = body[a - 1:b]; assert ssr == ssr.lower() and body[:a - 1] == body[:a - 1].upper(), c.id
    g5, c5, _ = run(os.path.join(d, "in"), os.path.join(d, "out5"), fasta_layout=("consolidated",), ssr_case="upper",
                    log=lambda *_: None)
    assert not os.path.exists(os.path.join(d, "out5", "fasta_per_locus"))
    rec = open(os.path.join(d, "out5", "SSRs_with_flanks.fasta")).read().split(">")[1].split("\n")
    body = "".join(rec[1:]); a, b = map(int, re.search(r"ssr_in_seq=(\d+)-(\d+)", rec[0]).groups())
    assert body[a - 1:b] == body[a - 1:b].upper() and body[:a - 1] == body[:a - 1].lower()   # SSR UPPERCASE
    # numbering follows the reference genome, top to bottom
    assert [c.id for c in clusters] == sorted(c.id for c in clusters)
    assert all(a.refpos <= b.refpos for a, b in zip(clusters, clusters[1:]))
    # zoom plot
    g3, c3, _ = run(os.path.join(d, "in"), os.path.join(d, "out3"), regions=[(1000, 8000)], log=lambda *_: None)
    assert os.path.exists(os.path.join(d, "out3", "micromine_region_1000-8000.svg"))
    # single-genome annotation mode (CLI accepts a file path) and 75 bp flanks
    single = os.path.join(d, "in", sorted(os.listdir(os.path.join(d, "in")))[0])
    g6, c6, ssrs6 = run(single, os.path.join(d, "out6"), flank=75, log=lambda *_: None)
    assert len(g6) == 1 and c6 == [] and len(ssrs6) >= 12
    out6 = os.path.join(d, "out6")
    assert all(os.path.exists(os.path.join(out6, f)) for f in ("micromine_report.xlsx", "micromine_SSRs.gff3", "SSRs_with_flanks.fasta", "micromine_annotation_map.svg"))
    assert all(len(s.left) == 75 or s.start <= 75 for s in ssrs6) and "flank=75/75" in open(os.path.join(out6, "SSRs_with_flanks.fasta")).read()
    import openpyxl
    wb6 = openpyxl.load_workbook(os.path.join(out6, "micromine_report.xlsx")); assert wb6.sheetnames[:3] == ["SSR_annotation", "Summary", "Motif_counts"]
    g7, c7, _ = run(os.path.join(d, "in"), os.path.join(d, "out7"), flank=75, log=lambda *_: None)   # comparison with 75 bp flanks
    assert all(len(m.left) <= 75 and len(m.right) <= 75 for c in c7 for m in c.members) and len(c7) == len(clusters)
    try:
        run(os.path.join(d, "in"), os.path.join(d, "out8"), flank=50, anchor=60, log=lambda *_: None); raise SystemExit("anchor > flank accepted")
    except ValueError:
        pass
    # N bases: SSRs whose flanking region has N are flagged (cross in plots, WARNING_N in FASTA/GFF3, Excel column)
    from micromine.core import read_fasta, mine_genome as _mg
    inn = os.path.join(d, "inN"); os.makedirs(inn)
    for f in os.listdir(os.path.join(d, "in")):
        cid, seq = read_fasta(os.path.join(d, "in", f))[0]
        if f.startswith("02_"):
            tgt = next(m for m in clusters[3].members if m.genome == 1)        # an SSR of species 02
            seq = seq[:tgt.start - 40] + "NNN" + seq[tgt.start - 37:]
        open(os.path.join(inn, f), "w").write(f">{cid}\n{seq}\n")
    gN, cN, sN = run(inn, os.path.join(d, "outN"), log=lambda *_: None)
    flagged = [m for c in cN for m in c.members if m.compromised]
    assert len(flagged) >= 1 and all(m.genome == 1 and m.n_flank == 3 for m in flagged), flagged
    assert "compromised_by_N_in=02" in open(os.path.join(d, "outN", "SSRs_with_flanks.fasta")).read()
    assert "WARNING_N=3" in open(os.path.join(d, "outN", "fasta_per_species", "02_snps_flanco_SSRs.fasta")).read()
    wbN = openpyxl.load_workbook(os.path.join(d, "outN", "micromine_report.xlsx"))
    assert any(r[18] == "02" for r in wbN["SSR_comparison"].iter_rows(min_row=2, values_only=True))
    assert "#c0392b" in open(os.path.join(d, "outN", "micromine_comparison.svg")).read()
    # characters other than A,C,G,T,N are rejected, listing every offending file
    bad = os.path.join(d, "bad"); os.makedirs(bad)
    for f in sorted(os.listdir(os.path.join(d, "in")))[:3]:
        txt = open(os.path.join(d, "in", f)).read()
        if f.startswith("02_") or f.startswith("03_"):
            ls = txt.split("\n"); ls[1] += "R"                      # first sequence line (not the header)
            txt = "\n".join(ls) if f.startswith("02_") else txt + "ACGU-\n"
        open(os.path.join(bad, f), "w").write(txt)
    try:
        run(bad, os.path.join(d, "outBad"), log=lambda *_: None); raise SystemExit("invalid bases accepted")
    except ValueError as e:
        assert "02_" in str(e) and "03_" in str(e) and "01_" not in str(e), str(e)
    print("OK:", len(clusters), "loci; rotation/reverse strand, flank SNPs, below-threshold search, numbering and zoom handled.")
