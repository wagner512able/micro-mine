"""Tabular (Excel) and FASTA outputs of MICRO-MINE."""
from __future__ import annotations

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import (CLASS_NAMES, VARIATION_TYPES, case_flanked, compromised_genomes, exclusivity, is_polymorphic,
                   relative_status, revcomp, variation_text, variation_type)

_FONT = Font(name="Arial", size=10)
_HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F3864")
_EXCL = {"real": "Real", "threshold": "Threshold-only"}


def _finish(ws, widths=None):
    for row in ws.iter_rows():
        for c in row:
            if c.row == 1:
                c.font, c.fill = _HEAD, _FILL
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                c.font = _FONT
    ws.freeze_panes = "A2"
    for i in range(1, ws.max_column + 1):
        w = (widths or {}).get(i)
        if w is None:
            vals = [len(str(ws.cell(r, i).value or "")) for r in range(1, min(ws.max_row, 200) + 1)]
            w = min(max(vals + [8]) + 2, 40)
        ws.column_dimensions[get_column_letter(i)].width = w


def _rep_label(c, genomes):
    parts = []
    for gi in sorted(set(c.genomes) | set(c.sub)):
        if gi in c.by_genome():
            reps = "/".join(str(m.reps) for m in c.by_genome()[gi])
            parts.append(f"{genomes[gi].tag}:({c.canon}){reps}")
        else:
            parts.append(f"{genomes[gi].tag}:[({c.canon}){c.sub[gi].reps}]")   # [ ] = below threshold
    return "; ".join(parts)


def _notes(c, n):
    notes = []
    kind = exclusivity(c, n)
    if kind == "threshold":
        notes.append("exclusive only because of the threshold: locus present with fewer repeats in other species")
    if c.cls == "mono" and (kind == "real" or is_polymorphic(c) or len(c.genomes) < n):
        notes.append("homopolymer: validate against reads/assembly before using as a marker")
    if compromised_genomes(c):
        notes.append("compromised: N bases in the SSR flanking region in at least one species")
    if max(len(v) for v in c.by_genome().values()) > 1:
        notes.append("multiple copies in the same species (possible IR/duplication)")
    return "; ".join(notes)


def _seq_header(ssr_case):
    return ("Sequence (flanks UPPERCASE, SSR lowercase)" if ssr_case != "upper"
            else "Sequence (flanks lowercase, SSR UPPERCASE)")


def write_xlsx(path, genomes, clusters, ssrs, params, ssr_case="lower"):
    wb = Workbook()
    n = len(genomes)

    ws = wb.active
    ws.title = "SSR_comparison"
    ws.append(["SSR_ID", "Motif (canonical)", "Class", "Reference species (first occurrence)", "Contig",
               "Start", "End", "Position on reference axis (bp)", "Present (>= threshold)",
               "Below threshold in", "No. of species", "Exclusive to", "Exclusivity type",
               "Repeats per species ([ ] = below threshold)", "Length polymorphism",
               "Variation type", "Differences vs. other species (longer/shorter/absent)",
               "Max copies per species", "Compromised by N in", "Notes", _seq_header(ssr_case)])
    for c in clusters:
        r = c.ref
        kind = exclusivity(c, n)
        ex = c.genomes[0] if kind else None
        ws.append([c.id, c.canon, c.cls, genomes[r.genome].label, r.contig, r.start, r.end,
                   round(c.refpos), ",".join(genomes[g].tag for g in c.genomes),
                   ",".join(genomes[g].tag for g in sorted(c.sub)), len(c.genomes),
                   genomes[ex].label if ex is not None else "", _EXCL.get(kind, ""),
                   _rep_label(c, genomes), "Yes" if is_polymorphic(c) else "No",
                   variation_type(c, n), variation_text(c, genomes),
                   max(len(v) for v in c.by_genome().values()),
                   ",".join(genomes[g].tag for g in compromised_genomes(c)), _notes(c, n),
                   case_flanked(r, ssr_case)])
    _finish(ws, {16: 30, 17: 55, 20: 50, 21: 60})

    wp = wb.create_sheet("Presence_Absence")
    wp.append(["SSR_ID", "Motif"] + [g.label for g in genomes])
    for c in clusters:
        present = set(c.genomes)
        wp.append([c.id, c.canon] + [1 if g.idx in present else ("sub" if g.idx in c.sub else 0)
                                      for g in genomes])
    wp.cell(wp.max_row + 2, 1, "1 = present (>= threshold); sub = locus present below threshold; "
                               "0 = no SSR at the locus")
    _finish(wp)

    wo = wb.create_sheet("SSR_per_species")
    wo.append(["SSR_ID", "File", "Species", "Contig", "Start", "End", "Observed motif", "Repeats",
               "Length (bp)", "Partial unit (highlighted, not counted)", "Orientation vs. reference", "Status",
               "N bases in flanks", _seq_header(ssr_case)])
    for c in clusters:
        rows = [(m, "meets threshold") for m in c.members] + \
               [(h, "below threshold") for h in c.sub.values()]
        for m, sit in sorted(rows, key=lambda x: (x[0].genome, x[0].contig, x[0].start)):
            g = genomes[m.genome]
            wo.append([c.id, g.tag, g.name, m.contig, m.start, m.end, m.motif, m.reps, len(m.seq), m.tail,
                       "+" if m.orient == 1 else "-", sit, m.n_flank, case_flanked(m, ssr_case)])
    _finish(wo, {14: 60})

    wr = wb.create_sheet("Summary")
    cls_cols = list(CLASS_NAMES.values())
    wr.append(["File", "Species", "Genome size (bp)"] + cls_cols
              + ["Total SSRs", "Real exclusive", "Threshold-only exclusive", "N bases in genome",
                 "SSRs compromised by N"])
    real = {g.idx: 0 for g in genomes}
    thr = {g.idx: 0 for g in genomes}
    for c in clusters:
        kind = exclusivity(c, n)
        if kind:
            (real if kind == "real" else thr)[c.genomes[0]] += 1
    for g in genomes:
        counts = {k: 0 for k in cls_cols}
        for s in ssrs:
            if s.genome == g.idx:
                counts[s.cls] += 1
        r = wr.max_row + 1
        wr.append([g.tag, g.name, g.length] + [counts[k] for k in cls_cols]
                  + [f"=SUM(D{r}:I{r})", real[g.idx], thr[g.idx], g.n_count,
                     sum(1 for s in ssrs if s.genome == g.idx and s.compromised)])
    _finish(wr)

    # variation of every locus in every species (longer / shorter / absent vs. the others)
    wv = wb.create_sheet("Variation_per_species")
    wv.append(["SSR_ID", "Motif (canonical)", "Class", "File", "Species", "Repeats",
               "Threshold status", "Relative to other species", "Median repeats of the other species"])
    rel_counts = {g.idx: {} for g in genomes}
    type_counts = {t: {k: 0 for k in CLASS_NAMES.values()} for t in VARIATION_TYPES}
    for c in clusters:
        type_counts[variation_type(c, n)][c.cls] += 1
        for gi, (rel, r, med, st) in relative_status(c, n).items():
            g = genomes[gi]
            wv.append([c.id, c.canon, c.cls, g.tag, g.name, r if r is not None else "",
                       {"meets": "meets threshold", "below": "below threshold"}.get(st, "no SSR"),
                       rel, med if med is not None else ""])
            rel_counts[gi][rel] = rel_counts[gi].get(rel, 0) + 1
            if st == "below":
                rel_counts[gi]["below"] = rel_counts[gi].get("below", 0) + 1
    _finish(wv)

    wt = wb.create_sheet("Variation_by_type")
    wt.append(["Variation type"] + list(CLASS_NAMES.values()) + ["Total"])
    for t in VARIATION_TYPES:
        r = wt.max_row + 1
        wt.append([t] + [type_counts[t][k] for k in CLASS_NAMES.values()] + [f"=SUM(B{r}:G{r})"])
    last = wt.max_row
    wt.append(["All loci"] + [f"=SUM({get_column_letter(i)}2:{get_column_letter(i)}{last})"
                              for i in range(2, 9)])
    _finish(wt, {1: 34})

    wb2 = wb.create_sheet("Variation_by_species")
    wb2.append(["File", "Species", "Longer than others", "Shorter than others", "Same as others",
                "Only species with SSR", "No SSR at locus", "Total loci",
                "Of which below threshold (subset of Shorter/Same)"])
    for g in genomes:
        rc = rel_counts[g.idx]
        r = wb2.max_row + 1
        wb2.append([g.tag, g.name, rc.get("Longer", 0), rc.get("Shorter", 0), rc.get("Same", 0),
                    rc.get("Only species with SSR at locus", 0), rc.get("No SSR at locus", 0),
                    f"=SUM(C{r}:G{r})", rc.get("below", 0)])
    _finish(wb2, {9: 30})

    wq = wb.create_sheet("Parameters")
    wq.append(["Parameter", "Value"])
    for k, v in params.items():
        wq.append([k, str(v)])
    _finish(wq, {1: 46, 2: 40})

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _comp_tag(c, genomes):
    g = compromised_genomes(c)
    return "|compromised_by_N_in=" + ",".join(genomes[i].tag for i in g) if g else ""


def _n_tag(m):
    return f"|WARNING_N={m.n_flank}" if m.n_flank else ""


def _partial_tag(m):
    return f"|partial={m.tail}" if m.tail else ""


def _pos_tag(left_len, region_len):
    """Position of the whole highlighted repeat tract inside the extracted sequence (1-based)."""
    return f"ssr_in_seq={left_len + 1}-{left_len + region_len}"


def _safe(name, limit=60):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(name))[:limit] or "contig"


def _wrap(seq, w=80):
    return "\n".join(seq[i:i + w] for i in range(0, len(seq), w))


def _rep_header(c, genomes):
    """Repeat number per species for FASTA headers (no spaces), e.g. 01:6,02:7,03:3(below_threshold)."""
    parts = []
    for gi in sorted(set(c.genomes) | set(c.sub)):
        if gi in c.by_genome():
            parts.append(f"{genomes[gi].tag}:" + "/".join(str(m.reps) for m in c.by_genome()[gi]))
        else:
            parts.append(f"{genomes[gi].tag}:{c.sub[gi].reps}(below_threshold)")
    return ",".join(parts)


def write_fastas(outdir, genomes, clusters, layouts=("species", "locus", "consolidated"),
                 ssr_case="lower"):
    """FASTA outputs. The repeat tract is highlighted by case (see case_flanked); every header carries the motif,
    the repeat number (`(AAG)6|repeats=6`) and the position of the tract inside the sequence (`ssr_in_seq`).

    species      : fasta_per_species/NN_species_SSRs.fasta - every SSR of one species/input file,
                   as found in its genome, with flanks
    locus        : fasta_per_locus/SSR.NN_class_motif.fasta - one file per SSR locus with the copy
                   from each species, reverse-complemented to the orientation of the first copy so
                   the records can be aligned directly (primer design)
    consolidated : SSRs_with_flanks.fasta - one record per locus (first species where it occurs)
    """
    layouts = set(layouts or ())
    if "species" in layouts:
        per = os.path.join(outdir, "fasta_per_species")
        os.makedirs(per, exist_ok=True)
        by_g: dict[int, list] = {g.idx: [] for g in genomes}
        for c in clusters:
            for m in c.members:
                by_g[m.genome].append(m)
        for g in genomes:
            with open(os.path.join(per, f"{g.label}_SSRs.fasta"), "w") as fh:
                for m in sorted(by_g[g.idx], key=lambda x: (x.contig, x.start)):
                    fh.write(f">{m.cluster}|{g.label}|{m.contig}:{m.start}-{m.end}|({m.motif}){m.reps}|repeats={m.reps}"
                             f"|flank={len(m.left)}/{len(m.right)}|{_pos_tag(len(m.left), len(m.region))}"
                             f"{_partial_tag(m)}{_n_tag(m)}\n{_wrap(case_flanked(m, ssr_case))}\n")
    if "locus" in layouts:
        per = os.path.join(outdir, "fasta_per_locus")
        os.makedirs(per, exist_ok=True)
        for c in clusters:
            with open(os.path.join(per, f"{c.id}_{c.cls}_{c.canon}.fasta"), "w") as fh:
                for m in sorted(c.members, key=lambda x: (x.genome, x.contig, x.start)):
                    seq, lleft, strand = case_flanked(m, ssr_case), len(m.left), "+"
                    if m.orient == -1:
                        seq, lleft, strand = revcomp(seq), len(m.right), "-"
                    fh.write(f">{c.id}|{genomes[m.genome].label}|{m.contig}:{m.start}-{m.end}|"
                             f"({c.canon}){m.reps}|repeats={m.reps}|strand={strand}|"
                             f"{_pos_tag(lleft, len(m.region))}{_partial_tag(m)}{_n_tag(m)}\n{_wrap(seq)}\n")
    if "consolidated" in layouts:
        with open(os.path.join(outdir, "SSRs_with_flanks.fasta"), "w") as fh:
            for c in clusters:
                r = c.ref
                pres = ",".join(genomes[i].tag for i in c.genomes)
                fh.write(f">{c.id}|({c.canon}){r.reps}|repeats={r.reps}|ref={genomes[r.genome].label}|"
                         f"{r.contig}:{r.start}-{r.end}|repeats_by_species={_rep_header(c, genomes)}|present_in={pres}|"
                         f"{_pos_tag(len(r.left), len(r.region))}{_partial_tag(r)}{_n_tag(r)}{_comp_tag(c, genomes)}\n"
                         f"{_wrap(case_flanked(r, ssr_case))}\n")


def write_exclusive_fasta(outdir, genomes, clusters, mode="species", single_copy=False,
                          include_threshold=False, ssr_case="lower"):
    """FASTA with ONLY the exclusive SSRs, ready for primer design (no manual filtering needed).

    A locus is exclusive when it meets the threshold in a single species. By default only REAL exclusive loci are
    exported (no homologous locus, even below the threshold, in the other species); threshold-only exclusives can
    be added. single_copy keeps only loci with one copy in that species (safer for primers).

    mode 'species': fasta_exclusive/NN_species_exclusive.fasta (one file per species)
    mode 'contig' : fasta_exclusive/NN_species__contig_exclusive.fasta (one file per species and contig)
    Returns (number of SSRs written, number of files).
    """
    if mode not in ("species", "contig"):
        return 0, 0
    n = len(genomes)
    groups: dict[str, list] = {}
    for c in clusters:
        kind = exclusivity(c, n)
        if kind is None or (kind == "threshold" and not include_threshold):
            continue
        gi = c.genomes[0]
        copies = c.by_genome()[gi]
        if single_copy and len(copies) > 1:
            continue
        for m in copies:
            key = genomes[gi].label if mode == "species" else f"{genomes[gi].label}__{_safe(m.contig)}"
            groups.setdefault(key, []).append((m, c, kind, len(copies), genomes[gi].label))
    if not groups:
        return 0, 0
    per = os.path.join(outdir, "fasta_exclusive")
    os.makedirs(per, exist_ok=True)
    total = 0
    for key, recs in sorted(groups.items()):
        with open(os.path.join(per, f"{key}_exclusive.fasta"), "w") as fh:
            for m, c, kind, nc, label in sorted(recs, key=lambda t: (t[0].contig, t[0].start)):
                fh.write(f">{c.id}|{label}|{m.contig}:{m.start}-{m.end}|({m.motif}){m.reps}|repeats={m.reps}|"
                         f"exclusive={'real' if kind == 'real' else 'threshold_only'}|locus_copies={nc}|"
                         f"flank={len(m.left)}/{len(m.right)}|{_pos_tag(len(m.left), len(m.region))}"
                         f"{_partial_tag(m)}{_n_tag(m)}\n{_wrap(case_flanked(m, ssr_case))}\n")
                total += 1
    return total, len(groups)
