"""Single-genome mode: SSR annotation of one FASTA file (no species comparison).

Outputs: Excel report, GFF3 annotation, FASTA of SSRs with flanks and a map of the SSRs along
the genome (SVG, PDF optional).
"""
from __future__ import annotations

import os
import urllib.parse
from collections import Counter

from openpyxl import Workbook

from . import __author__, __contact__, __version__
from .core import CLASS_NAMES, case_flanked, cluster_ssrs, mine_genome
from .plot import plot_annotation_svg, svg_to_pdf
from .report import _finish, _partial_tag, _n_tag, _pos_tag, _safe, _seq_header, _wrap


def number_ssrs(ssrs):
    """SSR.01, SSR.02... in genome order (contigs in file order, then position)."""
    width = max(2, len(str(len(ssrs))))
    for i, s in enumerate(ssrs, 1):
        s.cluster = f"SSR.{i:0{width}d}"
    return ssrs


def locus_copy_info(ssrs, anchor_len=60, sim_thr=0.6):
    """For every SSR: (copies of its locus in the genome, number of contigs carrying the locus).

    Loci are defined as in the comparison mode (same canonical motif + similar flanks, either strand), so SSRs in
    repeated sequences (e.g. inverted repeats or dispersed repeats) are recognised as multi-copy loci.
    """
    info = {}
    if ssrs:
        for c in cluster_ssrs(ssrs, sim_thr=sim_thr, anchor_len=anchor_len):
            n_contigs = len({m.contig for m in c.members})
            for m in c.members:
                info[id(m)] = (len(c.members), n_contigs)
    return info


def write_annotation_xlsx(path, genome, ssrs, mins, params, ssr_case="lower", info=None):
    from openpyxl.utils import get_column_letter  # noqa: F401  (kept local; used by _finish widths)
    wb = Workbook()
    ws = wb.active
    ws.title = "SSR_annotation"
    ws.append(["SSR_ID", "Contig", "Start", "End", "Length (bp)", "Class", "Motif (canonical)",
               "Observed motif", "Repeats", "Left flank (bp)", "Right flank (bp)", "N bases in flanks",
               "Compromised by N", "Locus copies in genome", "Contigs with this locus",
               "Partial unit (highlighted, not counted)", _seq_header(ssr_case)])
    for s in ssrs:
        ws.append([s.cluster, s.contig, s.start, s.end, len(s.seq), s.cls, s.canon, s.motif, s.reps,
                   len(s.left), len(s.right), s.n_flank, "Yes" if s.compromised else "No",
                   (info or {}).get(id(s), ("", ""))[0], (info or {}).get(id(s), ("", ""))[1], s.tail,
                   case_flanked(s, ssr_case)])
    _finish(ws, {17: 60})

    wsum = wb.create_sheet("Summary")
    wsum.append(["Class", "Min. repeats", "SSRs", "% of SSRs", "SSRs per kb"])
    by_class = Counter(s.cls for s in ssrs)
    for k, name in CLASS_NAMES.items():
        r = wsum.max_row + 1
        wsum.append([name, mins.get(k, ""), by_class.get(name, 0),
                     f"=IF($C$8=0,0,C{r}/$C$8)", f"=IF($B$10=0,0,C{r}/($B$10/1000))"])
    wsum.append(["Total", "", "=SUM(C2:C7)", "=IF($C$8=0,0,C8/$C$8)", "=IF($B$10=0,0,C8/($B$10/1000))"])
    wsum.cell(10, 1, "Genome size (bp)")
    wsum.cell(10, 2, genome.length)
    wsum.cell(11, 1, "Contigs")
    wsum.cell(11, 2, len(genome.contigs))
    wsum.cell(12, 1, "N bases in genome")
    wsum.cell(12, 2, genome.n_count)
    wsum.cell(13, 1, "SSRs compromised by N")
    wsum.cell(13, 2, sum(1 for s in ssrs if s.compromised))
    wsum.cell(14, 1, "SSRs in single-copy loci")
    wsum.cell(14, 2, sum(1 for s in ssrs if (info or {}).get(id(s), (0,))[0] == 1))
    _finish(wsum, {1: 18})
    for r in range(2, 9):
        wsum.cell(r, 4).number_format = "0.0%"
        wsum.cell(r, 5).number_format = "0.00"

    wm = wb.create_sheet("Motif_counts")
    wm.append(["Motif (canonical)", "Class", "SSRs", "Min. repeats observed", "Max. repeats observed",
               "Mean repeats"])
    motifs: dict[str, list] = {}
    for s in ssrs:
        motifs.setdefault(s.canon, []).append(s)
    for motif, ms in sorted(motifs.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        reps = [m.reps for m in ms]
        wm.append([motif, ms[0].cls, len(ms), min(reps), max(reps), round(sum(reps) / len(reps), 2)])
    _finish(wm)

    wq = wb.create_sheet("Parameters")
    wq.append(["Parameter", "Value"])
    for k, v in params.items():
        wq.append([k, str(v)])
    _finish(wq, {1: 46, 2: 40})
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def write_gff3(path, genome, ssrs, info=None):
    q = lambda t: urllib.parse.quote(str(t), safe="|:._-")
    with open(path, "w") as fh:
        fh.write("##gff-version 3\n")
        for cid, seq in genome.contigs:
            fh.write(f"##sequence-region {q(cid)} 1 {len(seq)}\n")
        for s in ssrs:
            attrs = (f"ID={s.cluster};Name=({s.motif}){s.reps};motif={s.motif};canonical_motif={s.canon};"
                     f"repeats={s.reps};class={s.cls};length={len(s.seq)}"
                     + (f";compromised_by_N={s.n_flank}" if s.n_flank else "")
                     + (f";partial_unit={s.tail}" if s.tail else "")
                     + (f";locus_copies={info[id(s)][0]}" if info and id(s) in info else ""))
            fh.write("\t".join([q(s.contig), "MicroMine", "microsatellite", str(s.start), str(s.end),
                                ".", "+", ".", attrs]) + "\n")


def _annotation_record(genome, s, ssr_case, info=None, extra=""):
    copies = f"|locus_copies={info[id(s)][0]}" if info and id(s) in info else ""
    return (f">{s.cluster}|{genome.label}|{s.contig}:{s.start}-{s.end}|({s.motif}){s.reps}|repeats={s.reps}{extra}{copies}"
            f"|flank={len(s.left)}/{len(s.right)}|{_pos_tag(len(s.left), len(s.region))}{_partial_tag(s)}{_n_tag(s)}\n"
            f"{_wrap(case_flanked(s, ssr_case))}\n")


def write_annotation_fasta(path, genome, ssrs, ssr_case="lower", info=None):
    with open(path, "w") as fh:
        for s in ssrs:
            fh.write(_annotation_record(genome, s, ssr_case, info))


def write_unique_fasta(outdir, genome, ssrs, info, mode="species", ssr_case="lower"):
    """FASTA with only the SSRs of single-copy loci (no similar locus anywhere else in the genome): the safest
    candidates for primer design. mode 'species': one file; mode 'contig': one file per contig.
    Returns (number of SSRs, number of files)."""
    keep = [s for s in ssrs if info.get(id(s), (0,))[0] == 1]
    if mode not in ("species", "contig") or not keep:
        return 0, 0
    if mode == "species":
        with open(os.path.join(outdir, "SSRs_unique_single_copy.fasta"), "w") as fh:
            for s in keep:
                fh.write(_annotation_record(genome, s, ssr_case, info))
        return len(keep), 1
    per = os.path.join(outdir, "fasta_unique")
    os.makedirs(per, exist_ok=True)
    by_contig: dict[str, list] = {}
    for s in keep:
        by_contig.setdefault(s.contig, []).append(s)
    for contig, lst in by_contig.items():
        with open(os.path.join(per, f"{genome.label}__{_safe(contig)}_unique.fasta"), "w") as fh:
            for s in lst:
                fh.write(_annotation_record(genome, s, ssr_case, info))
    return len(keep), len(by_contig)


def run_annotation(genome, output_dir, mins, flank=150, pdf=False, ssr_case="lower",
                   fasta_layout=("consolidated",), log=print, input_label=None, input_path=None,
                   anchor=60, similarity=0.6, exclusive_fasta="species"):
    log(f"[MICRO-MINE] single genome: annotation mode (no comparison) for {genome.label}")
    ssrs = mine_genome(genome, mins, flank)
    info = locus_copy_info(ssrs, min(anchor, flank), similarity)   # before numbering (clustering sets the ids)
    number_ssrs(ssrs)
    by_class = Counter(s.cls for s in ssrs)
    log(f"  {genome.label}: {genome.length:,} bp, {len(ssrs)} SSRs ("
        + ", ".join(f"{n} {by_class.get(n, 0)}" for n in CLASS_NAMES.values()) + ")")
    os.makedirs(output_dir, exist_ok=True)
    params = {"Software": "MICRO-MINE", "Version": __version__,
              "Developed by": f"{__author__} ({__contact__})", "Mode": "single-genome annotation",
              "Input": input_label or os.path.abspath(input_path or genome.path),
              "Flank length (bp)": flank, "Unique SSR FASTA": exclusive_fasta,
              "SSR highlight": "SSR lowercase, flanks UPPERCASE" if ssr_case != "upper" else "SSR UPPERCASE, flanks lowercase"}
    params.update({f"Min. repeats {name}": mins[k] for k, name in CLASS_NAMES.items()})
    write_annotation_xlsx(os.path.join(output_dir, "micromine_report.xlsx"), genome, ssrs, mins, params, ssr_case, info)
    write_gff3(os.path.join(output_dir, "micromine_SSRs.gff3"), genome, ssrs, info)
    if fasta_layout:
        write_annotation_fasta(os.path.join(output_dir, "SSRs_with_flanks.fasta"), genome, ssrs, ssr_case, info)
    if exclusive_fasta in ("species", "contig"):
        n_u, n_f = write_unique_fasta(output_dir, genome, ssrs, info, exclusive_fasta, ssr_case)
        log(f"  unique (single-copy) SSRs exported: {n_u} in {n_f} file(s)")
    svg = os.path.join(output_dir, "micromine_annotation_map.svg")
    plot_annotation_svg(svg, genome, ssrs)
    if pdf and not svg_to_pdf(svg, svg.replace(".svg", ".pdf")):
        log("[warning] PDF output requires 'pip install cairosvg'; the SVG file was created.")
    log(f"[MICRO-MINE] results in: {input_label and 'job output folder' or os.path.abspath(output_dir)}")
    return [genome], [], ssrs
